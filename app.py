import streamlit as st
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from datetime import datetime
import io

# ==================== KONFIGURACJA STRONY ====================
st.set_page_config(
    page_title="Kontrakt D/s | BDSM Contract Builder",
    page_icon="🔗",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==================== CUSTOM CSS – czytelny ciemny motyw ====================
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700&family=Inter:wght@400;500;600&display=swap');
    
    /* Główne tło */
    .stApp {
        background-color: #121212 !important;
        color: #f0f0f0 !important;
    }
    
    /* Nagłówki */
    h1, h2, h3, h4 {
        font-family: 'Playfair Display', Georgia, serif !important;
        color: #ffffff !important;
    }
    
    .main-title {
        font-family: 'Playfair Display', Georgia, serif;
        font-size: 2.5rem;
        font-weight: 700;
        color: #ffffff !important;
        text-align: center;
        margin-bottom: 0.25rem;
        letter-spacing: 0.5px;
    }
    
    .subtitle {
        text-align: center;
        color: #c0c0c0 !important;
        font-size: 1.05rem;
        margin-bottom: 1.8rem;
        font-weight: 400;
    }
    
    /* Pola tekstowe – białe tło + CZARNY tekst */
    .stTextInput > div > div > input,
    .stTextArea > div > div > textarea,
    .stTextInput input,
    .stTextArea textarea,
    div[data-testid="stTextInput"] input,
    div[data-testid="stTextArea"] textarea,
    input[type="text"],
    textarea {
        background-color: #ffffff !important;
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
        border: 1px solid #555555 !important;
        border-radius: 8px !important;
        font-size: 0.95rem !important;
        padding: 0.6rem 0.8rem !important;
        caret-color: #000000 !important;
    }
    
    .stTextInput > div > div > input::placeholder,
    .stTextArea > div > div > textarea::placeholder,
    input::placeholder,
    textarea::placeholder {
        color: #666666 !important;
        -webkit-text-fill-color: #666666 !important;
        opacity: 1 !important;
    }
    
    /* Selectbox */
    .stSelectbox > div > div,
    .stSelectbox [data-baseweb="select"],
    div[data-baseweb="select"] > div,
    div[data-baseweb="select"] span {
        background-color: #ffffff !important;
        color: #000000 !important;
        -webkit-text-fill-color: #000000 !important;
        border-radius: 8px !important;
    }
    
    .stSelectbox > div > div {
        border: 1px solid #555555 !important;
    }
    
    /* Etykiety pól */
    label, .stTextInput label, .stTextArea label, .stSelectbox label {
        color: #f0f0f0 !important;
        font-weight: 600 !important;
        font-size: 0.95rem !important;
    }
    
    /* Przyciski */
    .stButton > button {
        background-color: #8b2942 !important;
        color: #ffffff !important;
        border: none !important;
        border-radius: 8px !important;
        padding: 0.65rem 1.8rem !important;
        font-weight: 600 !important;
        font-size: 1rem !important;
        transition: all 0.2s ease !important;
    }
    
    .stButton > button:hover {
        background-color: #a3334f !important;
        box-shadow: 0 0 16px rgba(139, 41, 66, 0.5) !important;
    }
    
    .stDownloadButton > button {
        background-color: #2e7d4f !important;
        color: #ffffff !important;
        border-radius: 8px !important;
        font-weight: 600 !important;
        font-size: 1rem !important;
        padding: 0.7rem 1.8rem !important;
    }
    
    .stDownloadButton > button:hover {
        background-color: #3a9a62 !important;
        box-shadow: 0 0 16px rgba(46, 125, 79, 0.45) !important;
    }
    
    /* Boxy informacyjne */
    .info-box {
        background-color: #2a1f1f !important;
        border-left: 4px solid #c44d6a !important;
        padding: 1rem 1.2rem;
        border-radius: 0 8px 8px 0;
        margin: 1rem 0;
        font-size: 0.92rem;
        color: #f0e8e8 !important;
        line-height: 1.5;
    }
    
    .success-box {
        background-color: #1a2e22 !important;
        border-left: 4px solid #3a9a62 !important;
        padding: 1rem 1.2rem;
        border-radius: 0 8px 8px 0;
        margin: 1rem 0;
        color: #e0f0e8 !important;
        line-height: 1.5;
        font-size: 1.05rem;
    }
    
    /* Badges ról */
    .role-badge {
        display: inline-block;
        padding: 0.3rem 0.9rem;
        border-radius: 20px;
        font-size: 0.85rem;
        font-weight: 600;
        letter-spacing: 0.4px;
        margin-bottom: 0.6rem;
    }
    
    .badge-dom {
        background-color: #5c1e2e !important;
        color: #ffc0cb !important;
        border: 1px solid #c44d6a !important;
    }
    
    .badge-sub {
        background-color: #1e2a5c !important;
        color: #c0d0ff !important;
        border: 1px solid #4d6ac4 !important;
    }
    
    /* Sidebar – wymuszony ciemny motyw + wysoki kontrast */
    section[data-testid="stSidebar"],
    div[data-testid="stSidebar"] {
        background-color: #0f0f0f !important;
        border-right: 1px solid #333333 !important;
    }
    
    section[data-testid="stSidebar"] > div,
    div[data-testid="stSidebar"] > div {
        background-color: #0f0f0f !important;
    }
    
    /* Cały tekst w sidebarze */
    section[data-testid="stSidebar"] *,
    div[data-testid="stSidebar"] *,
    section[data-testid="stSidebar"] p,
    section[data-testid="stSidebar"] span,
    section[data-testid="stSidebar"] label,
    section[data-testid="stSidebar"] .stMarkdown,
    section[data-testid="stSidebar"] .stMarkdown p {
        color: #f0f0f0 !important;
    }
    
    /* Radio w sidebarze */
    section[data-testid="stSidebar"] .stRadio label {
        color: #ffffff !important;
        font-weight: 500 !important;
    }
    
    /* Info-box w sidebarze */
    section[data-testid="stSidebar"] .info-box,
    div[data-testid="stSidebar"] .info-box {
        background-color: #2a1a1a !important;
        border-left: 4px solid #c44d6a !important;
        color: #f5eaea !important;
    }
    
    section[data-testid="stSidebar"] .info-box b,
    div[data-testid="stSidebar"] .info-box b {
        color: #ffb0c0 !important;
    }
    
    /* Caption w sidebarze */
    section[data-testid="stSidebar"] .stCaption,
    section[data-testid="stSidebar"] small {
        color: #aaaaaa !important;
    }
    
    /* Nagłówki w sidebarze */
    section[data-testid="stSidebar"] h1,
    section[data-testid="stSidebar"] h2,
    section[data-testid="stSidebar"] h3,
    section[data-testid="stSidebar"] h4 {
        color: #ffffff !important;
    }
    
    /* Radio buttons */
    .stRadio > div {
        gap: 1.2rem;
    }
    
    .stRadio label {
        color: #f0f0f0 !important;
        font-size: 1rem !important;
    }
    
    /* Markdown i zwykły tekst */
    .stMarkdown, .stMarkdown p, .stMarkdown span {
        color: #e8e8e8 !important;
    }
    
    /* Separator */
    hr {
        border-color: #444444 !important;
        margin: 1.5rem 0 !important;
    }
    
    /* File uploader */
    .stFileUploader {
        background-color: #1e1e1e !important;
        border-radius: 8px;
        padding: 0.5rem;
    }
    
    /* Expander */
    .streamlit-expanderHeader {
        background-color: #1e1e1e !important;
        color: #ffffff !important;
        border-radius: 8px !important;
    }
    
    div[data-testid="stExpander"] {
        background-color: #1a1a1a !important;
        border: 1px solid #444444 !important;
        border-radius: 8px !important;
    }
    
    div[data-testid="stExpander"] p,
    div[data-testid="stExpander"] span,
    div[data-testid="stExpander"] pre,
    div[data-testid="stExpander"] code,
    div[data-testid="stExpander"] .stMarkdown {
        color: #f0f0f0 !important;
    }
    
    /* Caption */
    .stCaption, small {
        color: #aaaaaa !important;
    }
    
    /* Podgląd tekstu dokumentu */
    .preview-box {
        background-color: #1e1e1e !important;
        color: #f0f0f0 !important;
        border: 1px solid #444444 !important;
        border-radius: 8px !important;
        padding: 1.2rem !important;
        font-family: Georgia, serif !important;
        font-size: 0.9rem !important;
        line-height: 1.55 !important;
        white-space: pre-wrap !important;
        max-height: 400px !important;
        overflow-y: auto !important;
    }
</style>
""", unsafe_allow_html=True)

# ==================== FUNKCJE DOCX ====================

def set_run_font(run, name="Georgia", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_horizontal_line(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '8B3A3A')
    pBdr.append(bottom)
    pPr.append(pBdr)


def create_contract_docx(data: dict) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.3)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Georgia'
    style_normal.font.size = Pt(11)
    style_normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Georgia')

    # Tytuł
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("KONTRAKT")
    set_run_font(run, size=22, bold=True, color=(90, 30, 30))
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Dominujący / Dominująca  —  Uległy / Uległa")
    set_run_font(run, size=14, italic=True, color=(120, 60, 60))
    subtitle.paragraph_format.space_after = Pt(4)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_p.add_run(f"Data utworzenia: {datetime.now().strftime('%d.%m.%Y')}")
    set_run_font(run, size=9, color=(100, 100, 100))
    date_p.paragraph_format.space_after = Pt(12)

    add_horizontal_line(doc)

    # Preambuła
    h = doc.add_paragraph()
    run = h.add_run("1. Preambuła i zasady ogólne")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(6)

    preamble = (
        "Niniejszy dokument jest dobrowolnym, świadomym i wzajemnym porozumieniem pomiędzy stronami. "
        "Nie stanowi umowy prawnie wiążącej i nie zastępuje obowiązującego prawa. "
        "Każda ze stron ma prawo w każdej chwili, bez podawania przyczyny, wypowiedzieć ten kontrakt "
        "lub zawiesić jego obowiązywanie. Zasady SSC (Safe, Sane, Consensual) oraz RACK (Risk-Aware Consensual Kink) "
        "stanowią fundament tej relacji. Bezpieczeństwo fizyczne i psychiczne obu stron ma absolutny priorytet."
    )
    p = doc.add_paragraph()
    run = p.add_run(preamble)
    set_run_font(run, size=10)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15

    # Strony
    h = doc.add_paragraph()
    run = h.add_run("2. Strony kontraktu")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    run = p.add_run("Osoba Dominująca: ")
    set_run_font(run, size=11, bold=True)
    run = p.add_run(data.get("dom_name") or "……………………………………")
    set_run_font(run, size=11)

    p = doc.add_paragraph()
    run = p.add_run("Osoba Uległa: ")
    set_run_font(run, size=11, bold=True)
    run = p.add_run(data.get("sub_name") or "……………………………………")
    set_run_font(run, size=11)

    p = doc.add_paragraph()
    run = p.add_run("Zakres dynamiki: ")
    set_run_font(run, size=11, bold=True)
    run = p.add_run(data.get("scope") or "nie określono")
    set_run_font(run, size=11)
    p.paragraph_format.space_after = Pt(8)

    add_horizontal_line(doc)

    def add_section(title_text: str, content: str):
        h = doc.add_paragraph()
        run = h.add_run(title_text)
        set_run_font(run, size=12, bold=True, color=(90, 30, 30))
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)

        p = doc.add_paragraph()
        if content and content.strip():
            run = p.add_run(content.strip())
            set_run_font(run, size=10)
        else:
            run = p.add_run("[DO UZUPEŁNIENIA PRZEZ PARTNERA]")
            set_run_font(run, size=10, italic=True, color=(150, 80, 80))
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15

    # Część Dominującego
    h = doc.add_paragraph()
    run = h.add_run("3. Część Osoby Dominującej")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(4)

    add_section("3.1  Czego potrzebuję od osoby uległej", data.get("dom_needs", ""))
    add_section("3.2  Czego oczekuję", data.get("dom_expectations", ""))
    add_section("3.3  Hard Limits (kategorycznie nie)", data.get("dom_hard_limits", ""))
    add_section("3.4  Soft Limits (do omówienia / z warunkami)", data.get("dom_soft_limits", ""))
    add_section("3.5  Preferencje dotyczące aftercare i komunikacji", data.get("dom_aftercare", ""))

    add_horizontal_line(doc)

    # Część Uległego
    h = doc.add_paragraph()
    run = h.add_run("4. Część Osoby Uległej")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(4)

    add_section("4.1  Czego potrzebuję od osoby dominującej", data.get("sub_needs", ""))
    add_section("4.2  Czego oczekuję", data.get("sub_expectations", ""))
    add_section("4.3  Hard Limits (kategorycznie nie)", data.get("sub_hard_limits", ""))
    add_section("4.4  Soft Limits (do omówienia / z warunkami)", data.get("sub_soft_limits", ""))
    add_section("4.5  Preferencje dotyczące aftercare i komunikacji", data.get("sub_aftercare", ""))

    add_horizontal_line(doc)

    # Safewords
    h = doc.add_paragraph()
    run = h.add_run("5. Słowa bezpieczeństwa (Safewords)")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    run = p.add_run("System słów bezpieczeństwa: ")
    set_run_font(run, size=10, bold=True)
    run = p.add_run(data.get("safewords") or "RED – natychmiastowy stop  |  YELLOW – zwolnij / sprawdź  |  GREEN – wszystko w porządku")
    set_run_font(run, size=10)
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    run = p.add_run(
        "Użycie słowa bezpieczeństwa (szczególnie RED) nigdy nie jest powodem do kary ani dyskusji w momencie użycia. "
        "Scena zostaje natychmiast przerwana lub dostosowana."
    )
    set_run_font(run, size=9, italic=True, color=(80, 80, 80))
    p.paragraph_format.space_after = Pt(8)

    # Czas trwania
    h = doc.add_paragraph()
    run = h.add_run("6. Czas trwania i przeglądy")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    run = p.add_run(data.get("duration") or "Do ustalenia przez obie strony. Rekomendowany regularny przegląd (np. co 1–3 miesiące).")
    set_run_font(run, size=10)
    p.paragraph_format.space_after = Pt(10)

    add_horizontal_line(doc)

    # Potwierdzenie
    h = doc.add_paragraph()
    run = h.add_run("7. Potwierdzenie")
    set_run_font(run, size=13, bold=True, color=(90, 30, 30))
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    run = p.add_run(
        "Podpisując (lub potwierdzając) ten dokument, obie strony oświadczają, że zapoznały się z jego treścią, "
        "rozumieją go i dobrowolnie wyrażają zgodę na zawarte w nim ustalenia. "
        "Zgoda może zostać wycofana w dowolnym momencie."
    )
    set_run_font(run, size=9, italic=True)
    p.paragraph_format.space_after = Pt(16)

    table = doc.add_table(rows=2, cols=2)
    table.autofit = True

    cell = table.rows[0].cells[0]
    p = cell.paragraphs[0]
    run = p.add_run("Osoba Dominująca")
    set_run_font(run, size=10, bold=True)
    p = cell.add_paragraph()
    run = p.add_run(data.get("dom_name") or "________________________")
    set_run_font(run, size=10)
    p = cell.add_paragraph()
    run = p.add_run("Data: _______________")
    set_run_font(run, size=9, color=(100, 100, 100))

    cell = table.rows[0].cells[1]
    p = cell.paragraphs[0]
    run = p.add_run("Osoba Uległa")
    set_run_font(run, size=10, bold=True)
    p = cell.add_paragraph()
    run = p.add_run(data.get("sub_name") or "________________________")
    set_run_font(run, size=10)
    p = cell.add_paragraph()
    run = p.add_run("Data: _______________")
    set_run_font(run, size=9, color=(100, 100, 100))

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Dokument prywatny • Tylko do użytku stron • Nie stanowi umowy cywilnoprawnej")
    set_run_font(run, size=8, color=(130, 130, 130), italic=True)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def extract_text_from_docx(file) -> str:
    doc = Document(file)
    full = []
    for para in doc.paragraphs:
        if para.text.strip():
            full.append(para.text.strip())
    return "\n".join(full)


def parse_contract_data(text: str) -> dict:
    """Wyciąga dane podstawowe i sekcje z tekstu kontraktu."""
    import re
    data = {
        "dom_name": "",
        "sub_name": "",
        "scope": "",
        "duration": "",
        "safewords": "",
        "dom_needs": "",
        "dom_expectations": "",
        "dom_hard_limits": "",
        "dom_soft_limits": "",
        "dom_aftercare": "",
        "sub_needs": "",
        "sub_expectations": "",
        "sub_hard_limits": "",
        "sub_soft_limits": "",
        "sub_aftercare": "",
    }

    def grab(pattern, src):
        m = re.search(pattern, src, re.IGNORECASE | re.DOTALL)
        if m:
            val = m.group(1).strip()
            # Odrzuć placeholdery
            if val and "[DO UZUPEŁNIENIA" not in val and "……………………" not in val:
                return val
        return ""

    # Pola jednoliniowe – tylko do końca linii
    data["dom_name"] = grab(r"Osoba Dominująca:\s*([^\n]+)", text)
    data["sub_name"] = grab(r"Osoba Uległa:\s*([^\n]+)", text)
    data["scope"] = grab(r"Zakres dynamiki:\s*([^\n]+)", text)
    data["safewords"] = grab(r"System słów bezpieczeństwa:\s*([^\n]+)", text)

    # Czas trwania – sekcja 6 (jedna linia po nagłówku)
    data["duration"] = grab(
        r"6\.\s*Czas trwania i przeglądy\s*\n([^\n]+)", text
    )

    # Sekcje wieloliniowe – do następnego nagłówka sekcji
    data["dom_needs"] = grab(
        r"3\.1\s+Czego potrzebuję od osoby uległej\s*\n(.+?)(?=\n3\.2|\n4\.|\n5\.|$)", text
    )
    data["dom_expectations"] = grab(
        r"3\.2\s+Czego oczekuję\s*\n(.+?)(?=\n3\.3|\n4\.|\n5\.|$)", text
    )
    data["dom_hard_limits"] = grab(
        r"3\.3\s+Hard Limits[^\n]*\n(.+?)(?=\n3\.4|\n4\.|\n5\.|$)", text
    )
    data["dom_soft_limits"] = grab(
        r"3\.4\s+Soft Limits[^\n]*\n(.+?)(?=\n3\.5|\n4\.|\n5\.|$)", text
    )
    data["dom_aftercare"] = grab(
        r"3\.5\s+Preferencje dotyczące aftercare[^\n]*\n(.+?)(?=\n4\.|\n5\.|$)", text
    )

    data["sub_needs"] = grab(
        r"4\.1\s+Czego potrzebuję od osoby dominującej\s*\n(.+?)(?=\n4\.2|\n5\.|$)", text
    )
    data["sub_expectations"] = grab(
        r"4\.2\s+Czego oczekuję\s*\n(.+?)(?=\n4\.3|\n5\.|$)", text
    )
    data["sub_hard_limits"] = grab(
        r"4\.3\s+Hard Limits[^\n]*\n(.+?)(?=\n4\.4|\n5\.|$)", text
    )
    data["sub_soft_limits"] = grab(
        r"4\.4\s+Soft Limits[^\n]*\n(.+?)(?=\n4\.5|\n5\.|$)", text
    )
    data["sub_aftercare"] = grab(
        r"4\.5\s+Preferencje dotyczące aftercare[^\n]*\n(.+?)(?=\n5\.|$)", text
    )

    return data


# ==================== UI ====================

st.markdown('<div class="main-title">Kontrakt D/s</div>', unsafe_allow_html=True)
st.markdown(
    '<div class="subtitle">Generator kontraktu Dominujący/a — Uległy/a<br>Narzędzie komunikacji i bezpieczeństwa</div>',
    unsafe_allow_html=True
)

# Inicjalizacja session_state na plik
if "docx_bytes" not in st.session_state:
    st.session_state.docx_bytes = None
if "docx_filename" not in st.session_state:
    st.session_state.docx_filename = None

with st.sidebar:
    st.markdown("### Tryb pracy")
    mode = st.radio(
        "Wybierz:",
        ["Tworzę nowy kontrakt", "Uzupełniam otrzymany plik"],
        label_visibility="collapsed"
    )
    
    st.markdown("---")
    st.markdown("### Wskazówki")
    st.markdown("""
    <div class="info-box">
    <b>Hard Limits</b> – rzeczy, na które się kategorycznie nie zgadzasz. Nigdy nie powinny być negocjowane pod presją.<br><br>
    <b>Soft Limits</b> – obszary, które wymagają dodatkowej rozmowy, warunków lub wolniejszego wprowadzania.<br><br>
    Bądź konkretny/a. Im precyzyjniej opiszesz potrzeby i granice, tym bezpieczniejsza i przyjemniejsza będzie dynamika.
    </div>
    """, unsafe_allow_html=True)
    
    st.markdown("---")
    st.caption("Aplikacja działa lokalnie. Żadne dane nie są wysyłane na serwer.")


# ==================== TRYB 1: NOWY KONTRAKT ====================
if mode == "Tworzę nowy kontrakt":
    st.markdown("### Krok 1 — Kto wypełnia teraz?")
    filler = st.radio(
        "Wypełniam jako:",
        ["Osoba Dominująca", "Osoba Uległa"],
        horizontal=True
    )
    is_dom = filler == "Osoba Dominująca"
    
    st.markdown("---")
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown("#### Dane podstawowe")
        dom_name = st.text_input("Imię / pseudonim Osoby Dominującej", placeholder="np. Alex, Mistress K.")
        sub_name = st.text_input("Imię / pseudonim Osoby Uległej", placeholder="np. Sam, pet")
        scope = st.selectbox(
            "Zakres dynamiki",
            [
                "Tylko sceny (scene-only)",
                "Częściowy (part-time / weekendowy)",
                "24/7 z wyjątkami",
                "Total Power Exchange (TPE)",
                "Inny / do ustalenia"
            ]
        )
        duration = st.text_input(
            "Czas trwania / przeglądy",
            placeholder="np. 3 miesiące z przeglądem co miesiąc"
        )
        safewords = st.text_input(
            "Safewords",
            value="RED – stop  |  YELLOW – zwolnij  |  GREEN – ok"
        )
    
    with col2:
        st.markdown("#### Twoja część")
        if is_dom:
            st.markdown('<span class="role-badge badge-dom">DOMINUJĄCY / DOMINUJĄCA</span>', unsafe_allow_html=True)
            needs = st.text_area(
                "Czego potrzebuję od osoby uległej",
                height=110,
                placeholder="Np. szczerości, regularnej komunikacji, konkretnego rodzaju służby, zaufania..."
            )
            expectations = st.text_area(
                "Czego oczekuję",
                height=110,
                placeholder="Np. przestrzegania ustalonych reguł, inicjatywy w określonych obszarach..."
            )
            hard = st.text_area(
                "Hard Limits (kategorycznie nie)",
                height=100,
                placeholder="Np. blood play, publiczne ujawnianie tożsamości, konkretne praktyki..."
            )
            soft = st.text_area(
                "Soft Limits (do omówienia)",
                height=80,
                placeholder="Opcjonalnie..."
            )
            aftercare = st.text_area(
                "Aftercare i komunikacja – moje preferencje",
                height=80,
                placeholder="Np. potrzebuję ciszy po scenie / rozmowy / kontaktu fizycznego..."
            )
        else:
            st.markdown('<span class="role-badge badge-sub">ULEGŁY / ULEGŁA</span>', unsafe_allow_html=True)
            needs = st.text_area(
                "Czego potrzebuję od osoby dominującej",
                height=110,
                placeholder="Np. jasnych oczekiwań, poczucia bezpieczeństwa, konsekwencji, opieki po scenie..."
            )
            expectations = st.text_area(
                "Czego oczekuję",
                height=110,
                placeholder="Np. respektowania limitów, regularnego feedbacku, konkretnego stylu prowadzenia..."
            )
            hard = st.text_area(
                "Hard Limits (kategorycznie nie)",
                height=100,
                placeholder="Np. konkretne praktyki, słowa, sytuacje, które są absolutnie wykluczone..."
            )
            soft = st.text_area(
                "Soft Limits (do omówienia)",
                height=80,
                placeholder="Opcjonalnie..."
            )
            aftercare = st.text_area(
                "Aftercare i komunikacja – moje preferencje",
                height=80,
                placeholder="Np. potrzebuję przytulenia, rozmowy, wody, czasu w ciszy..."
            )
    
    st.markdown("---")
    
    generate = st.button("Generuj plik Word (.docx)", use_container_width=True, type="primary")
    
    if generate:
        data = {
            "dom_name": dom_name,
            "sub_name": sub_name,
            "scope": scope,
            "duration": duration,
            "safewords": safewords,
        }
        
        if is_dom:
            data.update({
                "dom_needs": needs,
                "dom_expectations": expectations,
                "dom_hard_limits": hard,
                "dom_soft_limits": soft,
                "dom_aftercare": aftercare,
                "sub_needs": "",
                "sub_expectations": "",
                "sub_hard_limits": "",
                "sub_soft_limits": "",
                "sub_aftercare": "",
            })
        else:
            data.update({
                "sub_needs": needs,
                "sub_expectations": expectations,
                "sub_hard_limits": hard,
                "sub_soft_limits": soft,
                "sub_aftercare": aftercare,
                "dom_needs": "",
                "dom_expectations": "",
                "dom_hard_limits": "",
                "dom_soft_limits": "",
                "dom_aftercare": "",
            })
        
        try:
            docx_bytes = create_contract_docx(data)
            filename = f"kontrakt_Ds_{datetime.now().strftime('%Y%m%d_%H%M')}.docx"
            
            st.session_state.docx_bytes = docx_bytes
            st.session_state.docx_filename = filename
            
            st.markdown("""
            <div class="success-box">
            ✅ Plik wygenerowany pomyślnie!<br>
            Kliknij zielony przycisk poniżej, aby pobrać. Potem prześlij plik partnerowi/partnerce.
            </div>
            """, unsafe_allow_html=True)
        except Exception as e:
            st.error(f"Błąd podczas generowania pliku: {e}")
    
    # Przycisk pobierania – zawsze widoczny jeśli plik jest w session_state
    if st.session_state.docx_bytes is not None:
        st.download_button(
            label="⬇️  Pobierz kontrakt (.docx)",
            data=st.session_state.docx_bytes,
            file_name=st.session_state.docx_filename or "kontrakt_Ds.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )


# ==================== TRYB 2: UZUPEŁNIANIE ====================
else:
    st.markdown("### Wgraj otrzymany plik kontraktu")
    
    uploaded = st.file_uploader(
        "Wybierz plik .docx wygenerowany wcześniej",
        type=["docx"],
        help="Wgraj plik, który dostałeś/aś od partnera/partnerki"
    )
    
    if uploaded:
        # Parsuj dokument
        try:
            uploaded.seek(0)
            text = extract_text_from_docx(uploaded)
            parsed = parse_contract_data(text)
        except Exception as e:
            text = ""
            parsed = {}
            st.error(f"Nie udało się odczytać pliku: {e}")

        # Wymuś wczytanie danych do session_state (omija cache widgetów Streamlit)
        file_id = getattr(uploaded, "name", "") + str(getattr(uploaded, "size", 0))
        if st.session_state.get("_parsed_file_id") != file_id:
            st.session_state["dn2"] = parsed.get("dom_name", "") or ""
            st.session_state["sn2"] = parsed.get("sub_name", "") or ""
            st.session_state["du2"] = parsed.get("duration", "") or ""
            st.session_state["sw2"] = parsed.get("safewords") or "RED – stop  |  YELLOW – zwolnij  |  GREEN – ok"
            st.session_state["_parsed_file_id"] = file_id
            st.session_state["_parsed_data"] = parsed
            # partner fields – ustawimy po wyborze roli, na razie wyczyść
            for k in ("on", "oe", "oh", "os", "oa"):
                if k in st.session_state:
                    del st.session_state[k]

        parsed = st.session_state.get("_parsed_data", parsed)

        st.markdown("""
        <div class="info-box">
        Plik wczytany. Dane podstawowe zostały automatycznie uzupełnione z dokumentu.
        Uzupełnij swoją część poniżej.
        </div>
        """, unsafe_allow_html=True)
        
        with st.expander("Podgląd treści wgranego pliku", expanded=False):
            if text:
                preview = text[:4000] + ("..." if len(text) > 4000 else "")
                preview_safe = (
                    preview.replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;")
                )
                st.markdown(
                    f'<div class="preview-box">{preview_safe}</div>',
                    unsafe_allow_html=True
                )
        
        st.markdown("---")
        st.markdown("### Uzupełnij swoją część")
        
        filler2 = st.radio(
            "Wypełniam jako:",
            ["Osoba Dominująca", "Osoba Uległa"],
            horizontal=True,
            key="filler2"
        )
        is_dom2 = filler2 == "Osoba Dominująca"
        
        # Opcje zakresu dynamiki
        scope_options = [
            "Tylko sceny (scene-only)",
            "Częściowy (part-time / weekendowy)",
            "24/7 z wyjątkami",
            "Total Power Exchange (TPE)",
            "Inny / do ustalenia"
        ]
        parsed_scope = parsed.get("scope", "")
        scope_index = 0
        for i, opt in enumerate(scope_options):
            if parsed_scope and (parsed_scope in opt or opt in parsed_scope):
                scope_index = i
                break
        
        col1, col2 = st.columns(2)
        
        with col1:
            st.markdown("#### Dane podstawowe *(zaciągnięte z pliku)*")
            dom_name2 = st.text_input(
                "Imię / pseudonim Osoby Dominującej",
                key="dn2"
            )
            sub_name2 = st.text_input(
                "Imię / pseudonim Osoby Uległej",
                key="sn2"
            )
            scope2 = st.selectbox(
                "Zakres dynamiki",
                scope_options,
                index=scope_index,
                key="sc2"
            )
            duration2 = st.text_input(
                "Czas trwania / przeglądy",
                key="du2"
            )
            safewords2 = st.text_input(
                "Safewords",
                key="sw2"
            )
        
        with col2:
            st.markdown("#### Twoja część")
            if is_dom2:
                st.markdown('<span class="role-badge badge-dom">DOMINUJĄCY / DOMINUJĄCA</span>', unsafe_allow_html=True)
                needs2 = st.text_area("Czego potrzebuję od osoby uległej", height=110, key="n2")
                expectations2 = st.text_area("Czego oczekuję", height=110, key="e2")
                hard2 = st.text_area("Hard Limits (kategorycznie nie)", height=100, key="h2")
                soft2 = st.text_area("Soft Limits (do omówienia)", height=80, key="s2")
                aftercare2 = st.text_area("Aftercare i komunikacja – moje preferencje", height=80, key="a2")
            else:
                st.markdown('<span class="role-badge badge-sub">ULEGŁY / ULEGŁA</span>', unsafe_allow_html=True)
                needs2 = st.text_area("Czego potrzebuję od osoby dominującej", height=110, key="n2b")
                expectations2 = st.text_area("Czego oczekuję", height=110, key="e2b")
                hard2 = st.text_area("Hard Limits (kategorycznie nie)", height=100, key="h2b")
                soft2 = st.text_area("Soft Limits (do omówienia)", height=80, key="s2b")
                aftercare2 = st.text_area("Aftercare i komunikacja – moje preferencje", height=80, key="a2b")
        
        st.markdown("---")
        st.markdown("#### Dane partnera *(zaciągnięte z pliku – możesz poprawić)*")
        st.caption("Te pola zostały automatycznie wypełnione z wgranego dokumentu. Dzięki temu finalny kontrakt będzie kompletny.")
        
        # Ustaw dane partnera w session_state przy zmianie roli / pliku
        role_key = f"{file_id}_{is_dom2}"
        if st.session_state.get("_partner_role_key") != role_key:
            if is_dom2:
                st.session_state["on"] = parsed.get("sub_needs", "") or ""
                st.session_state["oe"] = parsed.get("sub_expectations", "") or ""
                st.session_state["oh"] = parsed.get("sub_hard_limits", "") or ""
                st.session_state["os"] = parsed.get("sub_soft_limits", "") or ""
                st.session_state["oa"] = parsed.get("sub_aftercare", "") or ""
            else:
                st.session_state["on"] = parsed.get("dom_needs", "") or ""
                st.session_state["oe"] = parsed.get("dom_expectations", "") or ""
                st.session_state["oh"] = parsed.get("dom_hard_limits", "") or ""
                st.session_state["os"] = parsed.get("dom_soft_limits", "") or ""
                st.session_state["oa"] = parsed.get("dom_aftercare", "") or ""
            st.session_state["_partner_role_key"] = role_key
        
        if is_dom2:
            other_needs = st.text_area("Czego potrzebuje osoba uległa", height=70, key="on")
            other_exp = st.text_area("Czego oczekuje osoba uległa", height=70, key="oe")
            other_hard = st.text_area("Hard Limits osoby uległej", height=70, key="oh")
            other_soft = st.text_area("Soft Limits osoby uległej", height=60, key="os")
            other_after = st.text_area("Aftercare osoby uległej", height=60, key="oa")
        else:
            other_needs = st.text_area("Czego potrzebuje osoba dominująca", height=70, key="on")
            other_exp = st.text_area("Czego oczekuje osoba dominująca", height=70, key="oe")
            other_hard = st.text_area("Hard Limits osoby dominującej", height=70, key="oh")
            other_soft = st.text_area("Soft Limits osoby dominującej", height=60, key="os")
            other_after = st.text_area("Aftercare osoby dominującej", height=60, key="oa")
        
        generate2 = st.button("Generuj kompletny kontrakt (.docx)", use_container_width=True, type="primary", key="gen2")
        
        if generate2:
            data = {
                "dom_name": dom_name2,
                "sub_name": sub_name2,
                "scope": scope2,
                "duration": duration2,
                "safewords": safewords2,
            }
            
            if is_dom2:
                data.update({
                    "dom_needs": needs2,
                    "dom_expectations": expectations2,
                    "dom_hard_limits": hard2,
                    "dom_soft_limits": soft2,
                    "dom_aftercare": aftercare2,
                    "sub_needs": other_needs,
                    "sub_expectations": other_exp,
                    "sub_hard_limits": other_hard,
                    "sub_soft_limits": other_soft,
                    "sub_aftercare": other_after,
                })
            else:
                data.update({
                    "sub_needs": needs2,
                    "sub_expectations": expectations2,
                    "sub_hard_limits": hard2,
                    "sub_soft_limits": soft2,
                    "sub_aftercare": aftercare2,
                    "dom_needs": other_needs,
                    "dom_expectations": other_exp,
                    "dom_hard_limits": other_hard,
                    "dom_soft_limits": other_soft,
                    "dom_aftercare": other_after,
                })
            
            try:
                docx_bytes = create_contract_docx(data)
                filename = f"kontrakt_Ds_kompletny_{datetime.now().strftime('%Y%m%d_%H%M')}.docx"
                
                st.session_state.docx_bytes = docx_bytes
                st.session_state.docx_filename = filename
                
                st.markdown("""
                <div class="success-box">
                ✅ Kompletny kontrakt wygenerowany!<br>
                Kliknij zielony przycisk poniżej, aby pobrać.
                </div>
                """, unsafe_allow_html=True)
            except Exception as e:
                st.error(f"Błąd podczas generowania pliku: {e}")
        
        if st.session_state.docx_bytes is not None:
            st.download_button(
                label="⬇️  Pobierz kompletny kontrakt (.docx)",
                data=st.session_state.docx_bytes,
                file_name=st.session_state.docx_filename or "kontrakt_Ds_kompletny.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
                key="dl2"
            )

# Stopka
st.markdown("---")
st.markdown(
    "<div style='text-align:center; color:#888; font-size:0.8rem; padding:1rem 0;'>"
    "Narzędzie edukacyjne i komunikacyjne • Nie stanowi porady prawnej • Zawsze stawiaj bezpieczeństwo i zgodę na pierwszym miejscu"
    "</div>",
    unsafe_allow_html=True
)
